"""Compila imágenes en un documento Word, una imagen por hoja carta.

Reglas de diseño (no cambiar sin probar en Word):
- Hoja carta 8.5" x 11", márgenes laterales de 0.5" (ancho útil 7.5").
- Márgenes verticales fijos de 0.75"; la imagen se centra dentro de un
  bloque de 9.5" de alto desplazando solo el margen superior.
- Imágenes anchas se reducen al ancho útil; imágenes altas, al alto máximo
  de 9". Así siempre queda holgura para el salto de sección que crea Word.
- Interlineado sencillo y sin espacio antes/después: la altura del párrafo
  es la de la imagen.

La versión de escritorio (desktop-app/AIRIS_Converter.py) replica estas reglas.
"""
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image as PILImage

PAGE_W_IN = 8.5
PAGE_H_IN = 11.0
SIDE_MARGIN_IN = 0.5
TEXT_WIDTH_IN = 7.5        # PAGE_W_IN - 2 * SIDE_MARGIN_IN
VERTICAL_MARGIN_IN = 0.75
BOX_H_IN = 9.5             # bloque en el que se centra la imagen
MAX_IMAGE_H_IN = 9.0       # deja >= 0.25" de holgura para el salto de sección
SCREEN_DPI = 96            # píxeles -> pulgadas para dimensionar la imagen


def natural_key(path):
    """Clave de orden natural: 'doc_p2' va antes que 'doc_p10'."""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r'(\d+)', path)]


def image_size_in(img_path):
    """Devuelve (ancho, alto) de la imagen en pulgadas al escalar a 96 DPI."""
    with PILImage.open(img_path) as img:
        w_px, h_px = img.size
    return w_px / SCREEN_DPI, h_px / SCREEN_DPI


def fit_scale(img_w_in, img_h_in):
    """Factor de escala que nunca agranda y respeta ancho y alto máximos."""
    return min(1.0, TEXT_WIDTH_IN / img_w_in, MAX_IMAGE_H_IN / img_h_in)


def _configure_section(section, final_h):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(PAGE_W_IN)
    section.page_height = Inches(PAGE_H_IN)
    section.top_margin = Inches(VERTICAL_MARGIN_IN + (BOX_H_IN - final_h) / 2)
    section.bottom_margin = Inches(VERTICAL_MARGIN_IN)
    section.left_margin = section.right_margin = Inches(SIDE_MARGIN_IN)


def generate_word(images, output_path):
    """Crea el .docx con una imagen por hoja. Devuelve True si tuvo éxito."""
    try:
        doc = Document()
        images = sorted(images, key=natural_key)
        for i, img_path in enumerate(images):
            section = doc.add_section() if i > 0 else doc.sections[0]

            img_w_in, img_h_in = image_size_in(img_path)
            scale = fit_scale(img_w_in, img_h_in)
            final_w = img_w_in * scale
            final_h = img_h_in * scale
            _configure_section(section, final_h)

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.add_run().add_picture(img_path, width=Inches(final_w))
        doc.save(output_path)
        return True
    except Exception as e:
        print(f"Error en generate_word: {e}")
        return False
